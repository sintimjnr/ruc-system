import argparse
import os
import shutil
import subprocess
import sys
import zipfile
from datetime import datetime
from pathlib import Path

from backup_ruc import (
    EXPECTED_CATEGORIES,
    PG_RESTORE_FALLBACK,
    find_executable,
    load_dotenv_if_present,
    run_command,
)
from verify_backup import latest_backup_set, load_manifest, resolve_relative_file, verify_backup_set


BASE_DIR = Path(__file__).resolve().parent
RECOVERY_ROOT = Path(r"D:\RUC_SYSTEM_RECOVERY_TEST")
TEMP_DB_PREFIX = "ruc_system_recovery_test_"
PSQL_FALLBACK = r"C:\Program Files\PostgreSQL\13\bin\psql.exe"


class RestoreBlocked(RuntimeError):
    pass


class RestoreTestError(RuntimeError):
    pass


def now_timestamp():
    return datetime.now().strftime("%Y%m%d_%H%M%S")


def safe_child(parent, child):
    resolved_parent = parent.resolve()
    resolved_child = child.resolve()

    if resolved_child == resolved_parent or resolved_parent not in resolved_child.parents:
        raise RestoreTestError("Refusing to operate outside the recovery-test folder.")

    return resolved_child


def safe_extract_zip(zip_path, destination):
    destination = destination.resolve()
    destination.mkdir(parents=True, exist_ok=False)

    with zipfile.ZipFile(zip_path, "r") as archive:
        for member in archive.infolist():
            member_path = destination / member.filename
            safe_child(destination, member_path)

        archive.extractall(destination)


def validate_recovered_files(recovery_dir):
    checks = []

    for category in EXPECTED_CATEGORIES:
        category_path = recovery_dir / category
        exists = category_path.exists()
        checks.append((category, exists))

        if not exists:
            raise RestoreTestError(f"Recovered file tree is missing {category}.")

    workbook_checks = []

    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RestoreTestError("openpyxl is unavailable for workbook validation.") from exc

    candidates = []
    master_tracker = recovery_dir / "excel_files" / "master" / "NLZ_MASTER_TRACKER.xlsx"

    if master_tracker.exists():
        candidates.append(master_tracker)

    project_workbooks = sorted(
        path
        for path in (recovery_dir / "excel_files").glob("*.xlsx")
        if not path.name.startswith("~$")
    )

    if project_workbooks:
        candidates.append(project_workbooks[0])

    if not candidates:
        raise RestoreTestError("No representative XLSX files were recovered for validation.")

    for workbook in candidates:
        wb = load_workbook(workbook, read_only=True, data_only=False)
        workbook_checks.append((str(workbook.relative_to(recovery_dir)), len(wb.sheetnames)))
        wb.close()

    return {
        "categories": checks,
        "workbooks": workbook_checks,
    }


def admin_restore_context():
    load_dotenv_if_present()
    admin_user = os.environ.get("RUC_RESTORE_ADMIN_USER", "").strip()
    admin_password = os.environ.get("RUC_RESTORE_ADMIN_PASSWORD", "")
    runtime_user = os.environ.get("DB_USER", "").strip()

    if not admin_user or not admin_password:
        raise RestoreBlocked(
            "Temporary database restore was skipped because RUC_RESTORE_ADMIN_USER "
            "and RUC_RESTORE_ADMIN_PASSWORD are not configured."
        )

    if admin_user == "ruc_app" or (runtime_user and admin_user == runtime_user):
        raise RestoreBlocked(
            "Temporary database restore requires an administrative PostgreSQL role, not ruc_app."
        )

    host = os.environ.get("RUC_RESTORE_DB_HOST") or os.environ.get("DB_HOST") or "localhost"
    port = os.environ.get("RUC_RESTORE_DB_PORT") or os.environ.get("DB_PORT") or "5432"
    maintenance_db = os.environ.get("RUC_RESTORE_MAINTENANCE_DB") or "postgres"
    env = os.environ.copy()
    env["PGPASSWORD"] = admin_password

    return {
        "host": host,
        "port": port,
        "user": admin_user,
        "maintenance_db": maintenance_db,
        "env": env,
    }


def run_psql(context, database, sql):
    psql = find_executable("psql", "PSQL_EXE", PSQL_FALLBACK)
    command = [
        psql,
        "--no-password",
        "--set",
        "ON_ERROR_STOP=1",
        "--host",
        context["host"],
        "--port",
        str(context["port"]),
        "--username",
        context["user"],
        "--dbname",
        database,
        "--command",
        sql,
    ]
    return run_command(command, env=context["env"])


def create_temp_database(context, temp_db_name):
    if not temp_db_name.startswith(TEMP_DB_PREFIX) or temp_db_name == "ruc_system":
        raise RestoreTestError("Unsafe temporary database name.")

    result = run_psql(context, context["maintenance_db"], f'CREATE DATABASE "{temp_db_name}"')

    if result.returncode != 0:
        raise RestoreTestError("Temporary recovery database could not be created.")


def drop_temp_database(context, temp_db_name):
    if not temp_db_name.startswith(TEMP_DB_PREFIX) or temp_db_name == "ruc_system":
        raise RestoreTestError("Refusing to drop an unsafe database name.")

    terminate_sql = (
        "SELECT pg_terminate_backend(pid) FROM pg_stat_activity "
        f"WHERE datname = '{temp_db_name}' AND pid <> pg_backend_pid()"
    )
    run_psql(context, context["maintenance_db"], terminate_sql)
    result = run_psql(context, context["maintenance_db"], f'DROP DATABASE IF EXISTS "{temp_db_name}"')

    if result.returncode != 0:
        raise RestoreTestError("Temporary recovery database could not be dropped.")


def restore_temp_database(context, temp_db_name, database_backup_path):
    pg_restore = find_executable("pg_restore", "PG_RESTORE_EXE", PG_RESTORE_FALLBACK)
    command = [
        pg_restore,
        "--no-password",
        "--no-owner",
        "--host",
        context["host"],
        "--port",
        str(context["port"]),
        "--username",
        context["user"],
        "--dbname",
        temp_db_name,
        str(database_backup_path),
    ]
    result = run_command(command, env=context["env"])

    if result.returncode != 0:
        raise RestoreTestError("Temporary database restore failed.")


def validate_temp_database(context, temp_db_name):
    expected = ["admins", "projects", "employees", "telecom_sites", "audit_logs"]
    values = ", ".join("'" + table + "'" for table in expected)
    sql = (
        "SELECT table_name FROM information_schema.tables "
        "WHERE table_schema='public' AND table_name IN ("
        + values
        + ") ORDER BY table_name"
    )
    result = run_psql(context, temp_db_name, sql)

    if result.returncode != 0:
        raise RestoreTestError("Temporary database validation query failed.")

    found = set(result.stdout.split())
    missing = [table for table in expected if table not in found]

    if missing:
        raise RestoreTestError("Temporary database is missing tables: " + ", ".join(missing))

    return expected


def run_restore_test(backup_set, files_only=False, cleanup_files=False, cleanup_db=False):
    backup_set = Path(backup_set)
    verification = verify_backup_set(backup_set)

    if verification["status"] != "PASS":
        raise RestoreTestError("Backup verification failed before restore testing.")

    manifest = load_manifest(backup_set)
    archive_path = resolve_relative_file(backup_set, manifest["runtime_archive"])
    database_path = resolve_relative_file(backup_set, manifest["database"])
    timestamp = now_timestamp()
    recovery_dir = safe_child(RECOVERY_ROOT, RECOVERY_ROOT / timestamp)
    temp_db_name = f"{TEMP_DB_PREFIX}{timestamp.lower()}"
    result = {
        "backup_set": str(backup_set),
        "temporary_recovery_folder": str(recovery_dir),
        "temporary_database": temp_db_name,
        "file_restore": "NOT_RUN",
        "file_cleanup": "NOT_RUN",
        "database_restore": "NOT_RUN",
        "database_cleanup": "NOT_RUN",
        "recovered_workbooks": [],
        "recovered_tables": [],
    }

    try:
        safe_extract_zip(archive_path, recovery_dir)
        file_validation = validate_recovered_files(recovery_dir)
        result["file_restore"] = "PASS"
        result["recovered_workbooks"] = file_validation["workbooks"]

        if files_only:
            result["database_restore"] = "SKIPPED_FILES_ONLY"
        else:
            try:
                context = admin_restore_context()
                create_temp_database(context, temp_db_name)
                result["database_restore"] = "TEMP_DB_CREATED"
                restore_temp_database(context, temp_db_name, database_path)
                result["recovered_tables"] = validate_temp_database(context, temp_db_name)
                result["database_restore"] = "PASS"

                if cleanup_db:
                    drop_temp_database(context, temp_db_name)
                    result["database_cleanup"] = "PASS"
            except RestoreBlocked as exc:
                result["database_restore"] = "BLOCKED: " + str(exc)

        if cleanup_files:
            safe_child(RECOVERY_ROOT, recovery_dir)
            shutil.rmtree(recovery_dir)
            result["file_cleanup"] = "PASS"

        return result

    except Exception:
        if cleanup_files and recovery_dir.exists():
            safe_child(RECOVERY_ROOT, recovery_dir)
            shutil.rmtree(recovery_dir)
        raise


def main(argv=None):
    parser = argparse.ArgumentParser(
        description="Safely restore-test a RUC Stage 6 backup into temporary targets."
    )
    parser.add_argument(
        "backup_set",
        nargs="?",
        help="Path to a Stage 6 backup set. Defaults to the latest set.",
    )
    parser.add_argument(
        "--files-only",
        action="store_true",
        help="Extract and validate files only. Do not attempt temporary database restore.",
    )
    parser.add_argument(
        "--cleanup-files",
        action="store_true",
        help="Remove the specific temporary recovery folder after validation.",
    )
    parser.add_argument(
        "--cleanup-db",
        action="store_true",
        help="Drop the specific temporary recovery database after validation.",
    )
    args = parser.parse_args(argv)

    backup_set = Path(args.backup_set) if args.backup_set else latest_backup_set()

    if backup_set is None:
        print("RUC restore test FAILED")
        print("No Stage 6 backup sets were found.")
        return 1

    try:
        result = run_restore_test(
            backup_set,
            files_only=args.files_only,
            cleanup_files=args.cleanup_files,
            cleanup_db=args.cleanup_db,
        )
    except Exception as exc:
        print(f"RUC restore test FAILED: {exc}")
        return 1

    status = "PARTIAL" if str(result["database_restore"]).startswith("BLOCKED") else "PASS"
    print(f"RUC restore test {status}")
    print(f"Backup set: {result['backup_set']}")
    print(f"Temporary recovery folder: {result['temporary_recovery_folder']}")
    print(f"Temporary database: {result['temporary_database']}")
    print(f"File restore: {result['file_restore']}")
    print(f"File cleanup: {result['file_cleanup']}")
    print(f"Database restore: {result['database_restore']}")
    print(f"Database cleanup: {result['database_cleanup']}")

    for workbook, sheet_count in result["recovered_workbooks"]:
        print(f"Recovered workbook: {workbook} ({sheet_count} sheets)")

    if result["recovered_tables"]:
        print("Recovered database tables: " + ", ".join(result["recovered_tables"]))

    return 2 if status == "PARTIAL" else 0


if __name__ == "__main__":
    raise SystemExit(main())
