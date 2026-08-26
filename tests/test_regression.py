import os
import shutil
import tempfile
import time
import unittest
import zipfile
from datetime import date, datetime, timedelta
from io import BytesIO
from pathlib import Path
from unittest import mock

from openpyxl import Workbook, load_workbook

import app as ruc
import verify_backup


BASE_DIR = Path(__file__).resolve().parents[1]
STAGE6_BACKUP_SET = Path(
    os.environ.get(
        "RUC_STAGE6_BACKUP_SET",
        str(BASE_DIR / "backups" / "disaster" / "20260825_210820"),
    )
)


class RucRegressionTestCase(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls._audit_patch = mock.patch.object(ruc, "audit_event", lambda *args, **kwargs: None)
        cls._audit_patch.start()
        cls._tracker_download_patch = mock.patch.object(
            ruc,
            "record_master_tracker_download_for_edit",
            lambda *args, **kwargs: ruc.get_master_tracker_status(),
        )
        cls._tracker_download_patch.start()

    @classmethod
    def tearDownClass(cls):
        cls._tracker_download_patch.stop()
        cls._audit_patch.stop()

    def setUp(self):
        self.client = ruc.app.test_client()

    def login_as(self, role):
        accounts = {
            "super_admin": {
                "admin_id": 2,
                "admin": "Marian",
                "role": "super_admin",
                "employee_id": None,
            },
            "hr": {
                "admin_id": 81,
                "admin": "LU",
                "role": "hr",
                "employee_id": None,
            },
            "team_leader": {
                "admin_id": 82,
                "admin": "BENJAMIN SINTIM",
                "role": "team_leader",
                "employee_id": 17,
            },
        }

        with self.client.session_transaction() as session:
            session.clear()
            session.update(accounts[role])

    def csrf_for_public_form(self):
        response = self.client.get("/form/52143")
        self.assertEqual(response.status_code, 200)

        with self.client.session_transaction() as session:
            return session.get(ruc.CSRF_SESSION_KEY)

    def assert_security_headers(self, response):
        self.assertEqual(response.headers.get("X-Content-Type-Options"), "nosniff")
        self.assertEqual(response.headers.get("X-Frame-Options"), "SAMEORIGIN")
        self.assertEqual(
            response.headers.get("Referrer-Policy"),
            "strict-origin-when-cross-origin",
        )

    def get_status(self, path):
        response = self.client.get(path)

        try:
            return response.status_code
        finally:
            response.close()

    def post_status(self, path, data=None):
        response = self.client.post(path, data=data or {})

        try:
            return response.status_code
        finally:
            response.close()


class SecurityRegressionTests(RucRegressionTestCase):

    def test_app_import_and_route_count(self):
        self.assertEqual(len(ruc.app.url_map._rules), 88)

    def test_secret_key_is_required(self):
        with mock.patch.dict(os.environ, {"SECRET_KEY": ""}, clear=True):
            with self.assertRaises(ruc.ConfigurationError):
                ruc.get_secret_key()

    def test_healthz_methods(self):
        response = self.client.get("/healthz")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json(), {"status": "ok"})
        self.assertEqual(self.client.post("/healthz").status_code, 405)

    def test_login_page_and_csrf_rejection(self):
        response = self.client.get("/")
        self.assertEqual(response.status_code, 200)
        self.assert_security_headers(response)

        response = self.client.post("/", data={"username": "nobody", "password": "bad"})
        self.assertEqual(response.status_code, 400)

        with self.client.session_transaction() as session:
            session[ruc.CSRF_SESSION_KEY] = "expected-token"

        response = self.client.post(
            "/",
            data={
                "username": "nobody",
                "password": "bad",
                ruc.CSRF_FORM_FIELD: "wrong-token",
            },
        )
        self.assertEqual(response.status_code, 400)

    def test_public_form_post_requires_valid_csrf(self):
        response = self.client.post("/form/52143", data={})
        self.assertEqual(response.status_code, 400)

        with self.client.session_transaction() as session:
            session[ruc.CSRF_SESSION_KEY] = "expected-token"

        response = self.client.post(
            "/form/52143",
            data={ruc.CSRF_FORM_FIELD: "wrong-token"},
        )
        self.assertEqual(response.status_code, 400)

    def test_sensitive_routes_are_post_and_csrf_protected(self):
        self.login_as("super_admin")
        self.assertEqual(self.get_status("/logout"), 405)
        self.assertEqual(self.post_status("/logout"), 400)

        self.login_as("super_admin")
        self.assertEqual(self.get_status("/reset_system"), 405)
        self.assertEqual(self.post_status("/reset_system"), 400)

        self.login_as("super_admin")
        self.assertEqual(self.get_status("/reports/site/NL106/handover/package"), 405)
        self.assertEqual(
            self.post_status("/reports/site/NL106/handover/package"),
            400,
        )

    def test_security_configuration_defaults(self):
        self.assertFalse(ruc.app.debug)
        self.assertFalse(ruc.config_bool("RUC_TRUST_PROXY_HEADERS", False))
        self.assertTrue(ruc.app.config["SESSION_COOKIE_HTTPONLY"])
        self.assertEqual(ruc.app.config["SESSION_COOKIE_SAMESITE"], "Lax")
        self.assertEqual(ruc.LOGIN_RATE_LIMIT_ATTEMPTS, 5)
        self.assertEqual(ruc.LOGIN_RATE_LIMIT_WINDOW_SECONDS, 60)
        self.assertEqual(ruc.LOGIN_RATE_LIMIT_BLOCK_SECONDS, 300)

    def test_authenticated_html_is_not_cached(self):
        self.login_as("super_admin")
        response = self.client.get("/dashboard")
        self.assertEqual(response.status_code, 200)
        self.assertIn("no-store", response.headers.get("Cache-Control", ""))
        self.assertEqual(response.headers.get("Pragma"), "no-cache")
        self.assertEqual(response.headers.get("Expires"), "0")


class RbacRegressionTests(RucRegressionTestCase):

    def test_unauthenticated_protected_routes_redirect(self):
        for path in [
            "/dashboard",
            "/users",
            "/master_tracker",
            "/upload_master_tracker",
            "/reports",
        ]:
            with self.subTest(path=path):
                self.assertIn(self.get_status(path), {302, 401, 403})

    def test_super_admin_representative_access(self):
        self.login_as("super_admin")
        allowed_paths = [
            "/dashboard",
            "/users",
            "/projects",
            "/sites",
            "/reports",
            "/master_tracker",
            "/upload_master_tracker",
        ]

        for path in allowed_paths:
            with self.subTest(path=path):
                self.assertEqual(self.get_status(path), 200)

    def test_hr_scope_and_denials(self):
        self.login_as("hr")
        allowed_paths = ["/dashboard", "/search", "/safety", "/reports/personnel"]

        for path in allowed_paths:
            with self.subTest(path=path):
                self.assertEqual(self.get_status(path), 200)

        denied_paths = [
            "/master_tracker",
            "/upload_master_tracker",
            "/open_excel/52143",
            "/sites/new",
            "/daily_operations",
            "/telecom_tasks",
            "/site_assignments",
        ]

        for path in denied_paths:
            with self.subTest(path=path):
                self.assertEqual(self.get_status(path), 403)

    def test_team_leader_scope_and_denials(self):
        self.login_as("team_leader")
        allowed_paths = ["/team_leader", "/sites/NL105", "/sites/NL108"]

        for path in allowed_paths:
            with self.subTest(path=path):
                self.assertEqual(self.get_status(path), 200)

        denied_paths = [
            "/sites/NL106",
            "/master_tracker",
            "/upload_master_tracker",
            "/open_excel/52143",
            "/users",
            "/audit_logs",
            "/id_generator",
        ]

        for path in denied_paths:
            with self.subTest(path=path):
                self.assertEqual(self.get_status(path), 403)


class PublicFormRegressionTests(RucRegressionTestCase):

    def test_valid_project_form_get_issues_csrf(self):
        token = self.csrf_for_public_form()
        self.assertTrue(token)

    def test_upload_size_configuration_is_active(self):
        self.assertGreater(ruc.app.config["MAX_CONTENT_LENGTH"], 0)
        self.assertEqual(
            ruc.app.config["MAX_CONTENT_LENGTH"],
            int(os.environ.get("MAX_UPLOAD_BYTES", 16 * 1024 * 1024)),
        )

    def test_unsupported_upload_extension_fails_without_employee_insert(self):
        conn = ruc.connect_db()
        cursor = conn.cursor()
        cursor.execute("SELECT COUNT(*) FROM employees")
        before_count = cursor.fetchone()[0]
        cursor.close()
        conn.close()

        token = self.csrf_for_public_form()
        data = {
            ruc.CSRF_FORM_FIELD: token,
            "first_name": "Regression",
            "middle_name": "",
            "last_name": "Upload",
            "position": "Tester",
            "email": "regression@example.invalid",
            "mobile": "0000000000",
            "phone_type": "Test",
            "ftap_imei": "",
            "ftap_email": "",
            "philtower_imei": "",
            "philtower_email": "",
            "telecom_role": "Tester",
            "assigned_du_id": "NL106",
            "sec_number": "",
            "nbi_reference": "",
            "wah_reference": "",
            "first_aid_reference": "",
            "sec_expiry": "",
            "nbi_issue_date": "",
            "nbi_expiry_date": "",
            "wah_issue_date": "",
            "wah_expiry_date": "",
            "first_aid_issue_date": "",
            "first_aid_expiry_date": "",
            "photo": (BytesIO(b"not an image"), "bad.txt"),
            "nbi": (BytesIO(b"%PDF-1.4"), "nbi.pdf"),
            "certificate": (BytesIO(b"%PDF-1.4"), "cert.pdf"),
            "signature": (BytesIO(b"png"), "signature.png"),
        }
        response = self.client.post(
            "/form/52143",
            data=data,
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Unsupported file type", response.data)

        conn = ruc.connect_db()
        cursor = conn.cursor()
        cursor.execute("SELECT COUNT(*) FROM employees")
        after_count = cursor.fetchone()[0]
        cursor.close()
        conn.close()
        self.assertEqual(after_count, before_count)


class WorkbookSafetyRegressionTests(unittest.TestCase):

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.work_dir = Path(self.temp_dir.name)
        self.target = self.work_dir / "target.xlsx"
        self.candidate = self.work_dir / "candidate.xlsx"
        self.backup_dir = self.work_dir / "backups"
        self._create_workbook(self.target, "old")
        self._create_workbook(self.candidate, "new")

    def tearDown(self):
        self.temp_dir.cleanup()

    def _create_workbook(self, path, value):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Sheet"
        worksheet["A1"] = value
        workbook.save(path)

    def _cell_value(self, path):
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            return workbook["Sheet"]["A1"].value
        finally:
            workbook.close()

    def test_lock_acquisition_release_and_timeout(self):
        lock_path = Path(ruc.workbook_lock_path(str(self.target)))

        with ruc.workbook_write_lock(str(self.target), timeout=1):
            self.assertTrue(lock_path.exists())

            with self.assertRaises(ruc.WorkbookBusyError):
                with ruc.workbook_write_lock(str(self.target), timeout=0.01):
                    pass

        self.assertFalse(lock_path.exists())

    def test_stale_dead_lock_recovery(self):
        lock_path = Path(ruc.workbook_lock_path(str(self.target)))
        lock_path.parent.mkdir(parents=True, exist_ok=True)
        lock_path.write_text("pid=0\ncreated_at=2000-01-01T00:00:00\n", encoding="utf-8")
        old_time = time.time() - ruc.WORKBOOK_STALE_LOCK_SECONDS - 10
        os.utime(lock_path, (old_time, old_time))

        with ruc.workbook_write_lock(str(self.target), timeout=1):
            self.assertTrue(lock_path.exists())

        self.assertFalse(lock_path.exists())

    def test_atomic_replacement_success_and_backup(self):
        result = ruc.replace_persistent_workbook_file(
            str(self.target),
            str(self.candidate),
            expected_sheets=("Sheet",),
            backup_folder=str(self.backup_dir),
            operation="regression workbook replace",
        )

        self.assertEqual(self._cell_value(self.target), "new")
        self.assertTrue(result["backup_path"])
        self.assertTrue(Path(result["backup_path"]).exists())

    def test_invalid_candidate_leaves_original_valid(self):
        invalid_candidate = self.work_dir / "invalid.xlsx"
        invalid_candidate.write_bytes(b"not a workbook")

        with self.assertRaises(ruc.WorkbookSafetyError):
            ruc.replace_persistent_workbook_file(
                str(self.target),
                str(invalid_candidate),
                expected_sheets=("Sheet",),
                backup_folder=str(self.backup_dir),
                operation="regression invalid replace",
            )

        self.assertEqual(self._cell_value(self.target), "old")
        temp_files = list(self.work_dir.glob(ruc.WORKBOOK_TEMP_PREFIX + "*"))
        self.assertEqual(temp_files, [])

    def test_validation_failure_prevents_success_state(self):
        with self.assertRaises(ruc.WorkbookValidationError):
            ruc.validate_workbook_file(str(self.target), expected_sheets=("Missing",))

    def test_excel_safe_text_neutralizes_formula_like_business_text(self):
        cases = {
            "Benjamin": "Benjamin",
            "=1+1": "'=1+1",
            "+1+1": "'+1+1",
            "-1+1": "'-1+1",
            "@SUM(A1:A2)": "'@SUM(A1:A2)",
            "=-390532-21025-10": "'=-390532-21025-10",
        }

        for source_value, expected_value in cases.items():
            with self.subTest(source_value=source_value):
                self.assertEqual(ruc.excel_safe_text(source_value), expected_value)

    def test_excel_safe_row_preserves_numbers_dates_and_intentional_formulas(self):
        workbook_path = self.work_dir / "types.xlsx"
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Sheet"
        worksheet["A1"] = "=SUM(1,1)"
        ruc.append_excel_row(
            worksheet,
            [123, date(2026, 8, 26), "=1+1", "+1+1"],
            text_columns={3, 4},
        )
        workbook.save(workbook_path)
        workbook.close()

        reloaded = load_workbook(workbook_path, data_only=False)
        try:
            worksheet = reloaded["Sheet"]
            self.assertEqual(worksheet["A1"].data_type, "f")
            self.assertEqual(worksheet["A2"].value, 123)
            self.assertEqual(worksheet["B2"].value.date(), date(2026, 8, 26))
            self.assertEqual(worksheet["C2"].data_type, "s")
            self.assertEqual(worksheet["C2"].value, "'=1+1")
            self.assertEqual(worksheet["D2"].data_type, "s")
            self.assertEqual(worksheet["D2"].value, "'+1+1")
        finally:
            reloaded.close()

    def test_access_info_row_escapes_kojo_formula_like_mobile_on_temp_workbook(self):
        workbook_path = self.work_dir / "access_info.xlsx"
        workbook = ruc.build_project_workbook_template()
        ruc.write_access_info_row(
            workbook,
            {"region": "MINDANAO"},
            {
                "id": 1,
                "first_name": "KOJO",
                "middle_name": "",
                "last_name": "TETE",
                "position": "RIGGER",
                "mobile": "=-390532-21025-10",
                "email": "JUNIORSINTIMKOREE@GMAIL.COM",
                "phone_type": "ANDROID",
                "telecom_role": "RIGGER",
                "assigned_du_id": "NL106",
                "overall_safety_status": "MISSING",
            },
        )
        workbook.save(workbook_path)
        workbook.close()

        reloaded = load_workbook(workbook_path, data_only=False)
        try:
            worksheet = reloaded["ACCESS INFO"]
            self.assertEqual(worksheet["A2"].value, "KOJO TETE")
            self.assertEqual(worksheet["E2"].data_type, "s")
            self.assertEqual(worksheet["E2"].value, "'=-390532-21025-10")
            self.assertEqual(worksheet["M2"].value, "NL106")
        finally:
            reloaded.close()

    def test_report_workbook_escapes_text_without_changing_typed_values(self):
        workbook_path = self.work_dir / "report.xlsx"

        with ruc.app.test_request_context():
            workbook = ruc.build_report_workbook(
                "=Injected Title",
                [
                    (
                        "Formula Risk",
                        ["NAME", "COUNT", "DATE"],
                        [
                            ["=1+1", 42, date(2026, 8, 26)],
                            ["Benjamin", "+1+1", None],
                        ],
                    )
                ],
            )

        workbook.save(workbook_path)
        workbook.close()

        reloaded = load_workbook(workbook_path, data_only=False)
        try:
            summary = reloaded["SUMMARY"]
            worksheet = reloaded["Formula Risk"]
            self.assertEqual(summary["B1"].value, "'=Injected Title")
            self.assertEqual(worksheet["A2"].value, "'=1+1")
            self.assertEqual(worksheet["B2"].value, 42)
            self.assertEqual(worksheet["C2"].value.date(), date(2026, 8, 26))
            self.assertEqual(worksheet["A3"].value, "Benjamin")
            self.assertEqual(worksheet["B3"].value, "'+1+1")
        finally:
            reloaded.close()


class MasterTrackerWorkflowRegressionTests(unittest.TestCase):

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.work_dir = Path(self.temp_dir.name)
        self.target = self.work_dir / "NLZ_MASTER_TRACKER.xlsx"
        self.state_path = self.work_dir / "runtime" / "master_tracker_state.json"
        self.backup_dir = self.work_dir / "backups"
        self.original_safe_abs_path = ruc.safe_abs_path
        self._create_tracker(self.target, ["NL106", "NL106", "NL107"])

    def tearDown(self):
        self.temp_dir.cleanup()

    def _create_tracker(self, path, duids):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "GLOBE NLZ"
        worksheet.append(["DU ID", "Site Name", "TowerCo"])

        for index, duid in enumerate(duids, start=1):
            worksheet.append([duid, f"Site {index}", f"TowerCo {index}"])

        workbook.save(path)
        workbook.close()

    def _tracker_bytes(self, duids):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "GLOBE NLZ"
        worksheet.append(["DU ID", "Site Name", "TowerCo"])

        for index, duid in enumerate(duids, start=1):
            worksheet.append([duid, f"Uploaded Site {index}", f"Uploaded TowerCo {index}"])

        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()
        buffer.seek(0)
        return buffer

    def _login_as(self, client, role):
        accounts = {
            "super_admin": {
                "admin_id": 2,
                "admin": "Marian",
                "role": "super_admin",
                "employee_id": None,
            },
            "hr": {
                "admin_id": 81,
                "admin": "LU",
                "role": "hr",
                "employee_id": None,
            },
            "team_leader": {
                "admin_id": 82,
                "admin": "BENJAMIN SINTIM",
                "role": "team_leader",
                "employee_id": 17,
            },
        }

        with client.session_transaction() as session:
            session.clear()
            session.update(accounts[role])

    def _csrf(self, client):
        with client.session_transaction() as session:
            session[ruc.CSRF_SESSION_KEY] = "tracker-token"

        return "tracker-token"

    def _safe_abs_path_for_temp_uploads(self, *parts):
        if parts[:2] == ("excel_files", "master"):
            upload_dir = self.work_dir / "excel_files" / "master"
            upload_dir.mkdir(parents=True, exist_ok=True)
            return str(upload_dir / parts[-1])

        if parts[:2] == ("backups", "master"):
            self.backup_dir.mkdir(parents=True, exist_ok=True)
            return str(self.backup_dir)

        return self.original_safe_abs_path(*parts)

    def test_tracker_validation_preserves_existing_duplicates_and_rejects_new_duplicates(self):
        unchanged = self.work_dir / "unchanged.xlsx"
        new_unique = self.work_dir / "new_unique.xlsx"
        increased_duplicate = self.work_dir / "increased_duplicate.xlsx"
        new_duplicate = self.work_dir / "new_duplicate.xlsx"
        missing_existing = self.work_dir / "missing_existing.xlsx"
        malformed = self.work_dir / "malformed.xlsx"

        self._create_tracker(unchanged, ["NL106", "NL106", "NL107"])
        self._create_tracker(new_unique, ["NL106", "NL106", "NL107", "NL108"])
        self._create_tracker(increased_duplicate, ["NL106", "NL106", "NL107", "NL107"])
        self._create_tracker(new_duplicate, ["NL106", "NL106", "NL107", "NL108", "NL108"])
        self._create_tracker(missing_existing, ["NL106", "NL106", "NL108"])
        malformed.write_bytes(b"not an xlsx workbook")

        self.assertEqual(ruc.validate_tracker(str(self.target), str(unchanged)), "OK")
        self.assertEqual(ruc.validate_tracker(str(self.target), str(new_unique)), "OK")
        self.assertIn(
            "Duplicate DU ID detected: NL107",
            ruc.validate_tracker(str(self.target), str(increased_duplicate)),
        )
        self.assertIn(
            "Duplicate DU ID detected: NL108",
            ruc.validate_tracker(str(self.target), str(new_duplicate)),
        )
        self.assertIn(
            "Missing DU IDs detected",
            ruc.validate_tracker(str(self.target), str(missing_existing)),
        )

        with self.assertRaises(ruc.WorkbookSafetyError):
            ruc.replace_persistent_workbook_file(
                str(self.target),
                str(malformed),
                validator=ruc.validate_tracker,
                expected_sheets=("GLOBE NLZ",),
                backup_folder=str(self.backup_dir),
                operation="master tracker malformed regression",
            )

    def test_pending_state_survives_rejection_and_clears_after_valid_replacement(self):
        download_time = datetime(2026, 8, 26, 10, 0, 0)
        initial_checksum = ruc.file_sha256(self.target)

        initial_status = ruc.get_master_tracker_status(
            str(self.target), str(self.state_path), now=download_time
        )
        self.assertEqual(initial_status["status"], "CURRENT")

        pending_status = ruc.record_master_tracker_download_for_edit(
            admin_id=2,
            username="Marian",
            tracker_path=str(self.target),
            state_path=str(self.state_path),
            now=download_time,
        )
        self.assertEqual(pending_status["status"], "UPDATE PENDING")
        self.assertFalse(pending_status["is_overdue"])

        under_threshold = ruc.get_master_tracker_status(
            str(self.target),
            str(self.state_path),
            now=download_time + timedelta(hours=23),
        )
        self.assertEqual(under_threshold["status"], "UPDATE PENDING")
        self.assertFalse(under_threshold["is_overdue"])

        overdue = ruc.get_master_tracker_status(
            str(self.target),
            str(self.state_path),
            now=download_time + timedelta(hours=25),
        )
        self.assertEqual(overdue["status"], "UPDATE PENDING")
        self.assertTrue(overdue["is_overdue"])

        duplicate_candidate = self.work_dir / "duplicate_candidate.xlsx"
        self._create_tracker(duplicate_candidate, ["NL106", "NL106", "NL107", "NL107"])

        with self.assertRaises(ruc.WorkbookValidationError):
            ruc.replace_persistent_workbook_file(
                str(self.target),
                str(duplicate_candidate),
                validator=ruc.validate_tracker,
                expected_sheets=("GLOBE NLZ",),
                backup_folder=str(self.backup_dir),
                operation="master tracker duplicate regression",
            )

        self.assertEqual(ruc.file_sha256(self.target), initial_checksum)
        rejected_status = ruc.record_master_tracker_upload_rejected(
            "Duplicate DU ID detected: NL107",
            admin_id=2,
            username="Marian",
            tracker_path=str(self.target),
            state_path=str(self.state_path),
            now=download_time + timedelta(hours=1),
        )
        self.assertEqual(rejected_status["status"], "UPDATE PENDING")

        valid_candidate = self.work_dir / "valid_candidate.xlsx"
        self._create_tracker(valid_candidate, ["NL106", "NL106", "NL107", "NL108"])
        ruc.replace_persistent_workbook_file(
            str(self.target),
            str(valid_candidate),
            validator=ruc.validate_tracker,
            expected_sheets=("GLOBE NLZ",),
            backup_folder=str(self.backup_dir),
            operation="master tracker valid regression",
        )
        accepted_status = ruc.record_master_tracker_upload_accepted(
            admin_id=2,
            username="Marian",
            tracker_path=str(self.target),
            state_path=str(self.state_path),
            now=download_time + timedelta(hours=2),
        )

        self.assertEqual(accepted_status["status"], "CURRENT")
        self.assertFalse(accepted_status["is_pending"])
        self.assertNotEqual(accepted_status["approved_checksum"], initial_checksum)

    def test_missing_or_corrupt_pending_metadata_fails_safely(self):
        status = ruc.get_master_tracker_status(str(self.target), str(self.state_path))
        self.assertEqual(status["status"], "CURRENT")
        self.assertFalse(status["metadata_warning"])

        self.state_path.parent.mkdir(parents=True, exist_ok=True)
        self.state_path.write_text("{bad json", encoding="utf-8")

        status = ruc.get_master_tracker_status(str(self.target), str(self.state_path))
        self.assertEqual(status["status"], "CURRENT")
        self.assertTrue(status["metadata_warning"])

    def test_download_route_records_pending_only_for_authorized_super_admin(self):
        client = ruc.app.test_client()

        patches = [
            mock.patch.object(ruc, "MASTER_TRACKER_PATH", str(self.target)),
            mock.patch.object(ruc, "MASTER_TRACKER_STATE_PATH", str(self.state_path)),
            mock.patch.object(ruc, "audit_event", lambda *args, **kwargs: None),
        ]

        with patches[0], patches[1], patches[2]:
            self.assertIn(client.get("/master_tracker").status_code, {302, 401, 403})
            self.assertFalse(self.state_path.exists())

            self._login_as(client, "hr")
            self.assertEqual(client.get("/master_tracker").status_code, 403)
            self.assertFalse(self.state_path.exists())

            self._login_as(client, "team_leader")
            self.assertEqual(client.get("/master_tracker").status_code, 403)
            self.assertFalse(self.state_path.exists())

            self._login_as(client, "super_admin")
            response = client.get("/master_tracker")
            try:
                self.assertEqual(response.status_code, 200)
            finally:
                response.close()

            self.assertTrue(self.state_path.exists())
            status = ruc.get_master_tracker_status(str(self.target), str(self.state_path))
            self.assertEqual(status["status"], "UPDATE PENDING")

    def test_upload_route_rejected_candidate_keeps_pending_and_valid_upload_clears_it(self):
        client = ruc.app.test_client()
        self._login_as(client, "super_admin")
        token = self._csrf(client)
        audit_actions = []

        patches = [
            mock.patch.object(ruc, "MASTER_TRACKER_PATH", str(self.target)),
            mock.patch.object(ruc, "MASTER_TRACKER_STATE_PATH", str(self.state_path)),
            mock.patch.object(ruc, "safe_abs_path", self._safe_abs_path_for_temp_uploads),
            mock.patch.object(
                ruc,
                "audit_event",
                lambda action, *args, **kwargs: audit_actions.append(action),
            ),
        ]

        with patches[0], patches[1], patches[2], patches[3]:
            ruc.record_master_tracker_download_for_edit(
                admin_id=2,
                username="Marian",
                tracker_path=str(self.target),
                state_path=str(self.state_path),
                now=datetime(2026, 8, 26, 10, 0, 0),
            )

            rejected = client.post(
                "/upload_master_tracker",
                data={
                    ruc.CSRF_FORM_FIELD: token,
                    "tracker": (
                        self._tracker_bytes(["NL106", "NL106", "NL107", "NL107"]),
                        "tracker.xlsx",
                    ),
                },
                content_type="multipart/form-data",
            )
            self.assertEqual(rejected.status_code, 302)
            self.assertEqual(
                ruc.get_master_tracker_status(str(self.target), str(self.state_path))[
                    "status"
                ],
                "UPDATE PENDING",
            )
            self.assertIn("MASTER_TRACKER_UPLOAD_REJECTED", audit_actions)

            accepted = client.post(
                "/upload_master_tracker",
                data={
                    ruc.CSRF_FORM_FIELD: token,
                    "tracker": (
                        self._tracker_bytes(["NL106", "NL106", "NL107", "NL108"]),
                        "tracker.xlsx",
                    ),
                },
                content_type="multipart/form-data",
            )
            self.assertEqual(accepted.status_code, 302)
            self.assertEqual(
                ruc.get_master_tracker_status(str(self.target), str(self.state_path))[
                    "status"
                ],
                "CURRENT",
            )
            self.assertIn("MASTER_TRACKER_UPLOAD_ACCEPTED", audit_actions)


class BackupVerificationRegressionTests(unittest.TestCase):

    def test_stage6_backup_set_verifies(self):
        self.assertTrue(STAGE6_BACKUP_SET.exists(), f"Missing backup set: {STAGE6_BACKUP_SET}")
        result = verify_backup.verify_backup_set(STAGE6_BACKUP_SET)
        self.assertEqual(result["status"], "PASS")

    def test_stage6_backup_archive_contains_required_categories(self):
        manifest = verify_backup.load_manifest(STAGE6_BACKUP_SET)
        archive_path = verify_backup.resolve_relative_file(
            STAGE6_BACKUP_SET,
            manifest["runtime_archive"],
        )

        with zipfile.ZipFile(archive_path, "r") as archive:
            names = archive.namelist()

        for category in ["excel_files", "uploads", "static/uploads", "id_cards", "generated_reports"]:
            with self.subTest(category=category):
                self.assertTrue(
                    any(
                        name.rstrip("/") == category or name.startswith(category + "/")
                        for name in names
                    )
                )


if __name__ == "__main__":
    unittest.main(verbosity=2)
